"""Tests for the extraction logic. Plain functions with asserts: run with `pytest`."""

import io
import zipfile

import pymupdf
from docx import Document

import extraction
from extraction import ExtractionError, extract, extract_docx, extract_pdf


def make_pdf(pages: list[str | None]) -> bytes:
    """A PDF with one page per item: text, or None for an image-only (scanned-style) page."""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text is None:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 80, 80), False)
            pix.set_rect(pix.irect, (190, 190, 190))
            page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pix)
        else:
            page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def png_bytes() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.set_rect(pix.irect, (10, 120, 200))
    return pix.tobytes("png")


def make_docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def raises(code: str, func, *args):
    try:
        func(*args)
    except ExtractionError as exc:
        assert exc.code == code, f"expected {code}, got {exc.code}"
        assert exc.message and "Traceback" not in exc.message
        return
    raise AssertionError(f"expected ExtractionError({code})")


# ---- PDF ---------------------------------------------------------------------------------------

def test_pdf_text_is_extracted_per_page_with_counts():
    result = extract_pdf(make_pdf(["A loop repeats a block of code until a condition stops it.",
                                   "An array is an ordered collection of values."]))
    assert "A loop repeats a block of code" in result["text"]
    assert "Page 2:" in result["text"] and "ordered collection" in result["text"]
    assert [p["number"] for p in result["pages"]] == [1, 2]
    assert all(p["chars"] > 30 for p in result["pages"])
    assert result["warnings"] == []


def test_pdf_scanned_page_is_reported_not_hidden():
    result = extract_pdf(make_pdf(["Typed text on the first page, long enough to count.", None]))
    assert result["pages"][1]["chars"] == 0
    assert any("Page 2" in w and "no readable text" in w for w in result["warnings"])


def test_pdf_all_image_pages_report_zero_characters():
    result = extract_pdf(make_pdf([None, None, None]))
    assert [p["chars"] for p in result["pages"]] == [0, 0, 0]
    assert len(result["warnings"]) == 3


def test_pdf_corrupt_and_empty_input_is_refused_cleanly():
    raises("unreadable", extract_pdf, b"this is not a pdf")
    raises("unreadable", extract_pdf, b"")


def test_pdf_page_limit():
    original = extraction.MAX_PAGES
    extraction.MAX_PAGES = 2
    try:
        raises("too_many_pages", extract_pdf, make_pdf(["one", "two", "three"]))
    finally:
        extraction.MAX_PAGES = original


def test_pdf_encrypted_is_refused():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    doc.close()
    raises("encrypted", extract_pdf, data)


def test_text_is_cut_off_at_the_limit_with_a_warning():
    original = extraction.MAX_TEXT_CHARS
    extraction.MAX_TEXT_CHARS = 40
    try:
        result = extract_pdf(make_pdf(["x" * 200]))
    finally:
        extraction.MAX_TEXT_CHARS = original
    assert len(result["text"]) == 40
    assert any("cut off" in w for w in result["warnings"])


# ---- DOCX --------------------------------------------------------------------------------------

def test_docx_paragraphs_and_tables_are_read_in_document_order():
    def build(doc):
        doc.add_paragraph("First paragraph.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Loop"
        table.cell(0, 1).text = "Repeats"
        table.cell(1, 0).text = "Array"
        table.cell(1, 1).text = "Ordered values"
        doc.add_paragraph("Last paragraph.")

    result = extract_docx(make_docx(build))
    text = result["text"]
    assert text.index("First paragraph.") < text.index("Loop | Repeats") < text.index("Last paragraph.")
    assert "Array | Ordered values" in text
    assert result["pages"] == [{"number": 1, "chars": len(text)}]
    assert result["warnings"] == []


def test_docx_blank_paragraphs_are_skipped():
    result = extract_docx(make_docx(lambda d: (d.add_paragraph(""), d.add_paragraph("Only this."), d.add_paragraph("  "))))
    assert result["text"] == "Only this."


def test_docx_images_produce_a_warning():
    def build(doc):
        doc.add_paragraph("Answer is in the picture below.")
        doc.add_picture(io.BytesIO(png_bytes()))

    result = extract_docx(make_docx(build))
    assert any("1 image" in w for w in result["warnings"])


def test_docx_that_is_not_a_docx_is_refused():
    raises("unreadable", extract_docx, b"not a zip at all")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "just a zip")
    raises("unreadable", extract_docx, buffer.getvalue())


def test_docx_zip_bomb_guard():
    original = extraction.MAX_DOCX_UNCOMPRESSED_BYTES
    extraction.MAX_DOCX_UNCOMPRESSED_BYTES = 100
    try:
        raises("too_large", extract_docx, make_docx(lambda d: d.add_paragraph("hello " * 50)))
    finally:
        extraction.MAX_DOCX_UNCOMPRESSED_BYTES = original


# ---- dispatch ----------------------------------------------------------------------------------

def test_extract_dispatches_by_type_and_rejects_others():
    assert "hello" in extract(make_pdf(["hello world, this page is long enough"]), "pdf")["text"]
    raises("unsupported_type", extract, b"x", "exe")
